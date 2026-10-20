"""Tests sur un Mac neuf (CI) : conversions Office, recherche, roue et panneau de recherche.

    python mac_ci.py <dossier de résultats>
"""
import importlib.util as u, json, os, shutil, subprocess, sys, tempfile, time
from pathlib import Path

D = Path.home() / "Library" / "Application Support" / "FileForge"
PY = D / "venv" / "bin" / "python"
OUT = Path(sys.argv[1]); OUT.mkdir(parents=True, exist_ok=True)
CMD = Path(tempfile.gettempdir()) / "ff_cmd.js"; RES = CMD.parent / "ff_cmd.out"
fails = []


def expect(name, cond, info=""):
    print(("OK   " if cond else "FAIL ") + name, info, flush=True)
    if not cond:
        fails.append(name)


spec = u.spec_from_file_location("ff", D / "fileforge.py")
ff = u.module_from_spec(spec); spec.loader.exec_module(ff)
ff.notify = lambda m: None
ff.choose = lambda items, prompt, default=None, **k: items[0]
T = Path.home() / "Documents" / "Essais FileForge"
shutil.rmtree(T, ignore_errors=True); T.mkdir(parents=True)

# 1. conversions Office (LibreOffice récupéré par l'installateur)
expect("LibreOffice disponible", bool(ff.soffice()), ff.soffice())
from pptx import Presentation
pr = Presentation()
s = pr.slides.add_slide(pr.slide_layouts[1]); s.shapes.title.text = "Séance 1 — été"
pptx = T / "PPT Séance 1.pptx"; pr.save(pptx)
import openpyxl
wb = openpyxl.Workbook(); wb.active.append(["Café", 3.5]); xlsx = T / "Budget été.xlsx"; wb.save(xlsx)
import pymupdf
for f in (pptx, xlsx):
    try:
        out = ff.doc_to_pdf([str(f)])[0]
        expect(f"{f.suffix} → PDF", pymupdf.open(out).page_count > 0, Path(out).name)
    except Exception as e:
        expect(f"{f.suffix} → PDF", False, str(e)[:120])

# 2. recherche
sp = u.spec_from_file_location("s", D / "search.py"); s = u.module_from_spec(sp); sp.loader.exec_module(s)
(T / "Facture EDF été 2026.pdf").write_bytes(b"%PDF-1.4\n%%EOF\n")
n, dt = s.build_index()
res, info = s.search("factrue edf", 5)
expect("recherche tolérante (factrue edf)", any(r["name"] == "Facture EDF été 2026.pdf" for r in res), f"{n} éléments indexés en {dt:.1f} s")


def js(code, timeout=6):
    RES.unlink(missing_ok=True); CMD.write_text(code, encoding="utf-8")
    t = time.time()
    while time.time() - t < timeout:
        if RES.exists() and not CMD.exists():
            time.sleep(0.05); return RES.read_text(encoding="utf-8")
        time.sleep(0.05)
    CMD.unlink(missing_ok=True); return "TIMEOUT"


def launch(args, want):
    env = dict(os.environ, FF_DEBUG="1")
    t0 = time.time()
    p = subprocess.Popen([str(PY), str(D / "wheel.py"), *args], env=env)
    while time.time() - t0 < 30:
        if js("mode", timeout=0.5).strip() == want:
            return p, time.time() - t0
        time.sleep(0.05)
    return p, None


def shot(name):
    subprocess.run(["screencapture", "-x", str(OUT / f"mac-{name}.png")], capture_output=True)


# 3. la roue sur un fichier, la zone de dépôt, le panneau de recherche
for args, want, name in (([str(pptx)], "main", "roue"), (["--no-selection"], "drop", "depot"), (["--search"], "find", "recherche")):
    p, dt = launch(args, want)
    expect(f"{name} : affichée", dt is not None, f"{dt and round(dt, 2)} s")
    time.sleep(1); shot(name)
    js("api('__close__'); true", timeout=2)
    try:
        p.wait(8)
    except subprocess.TimeoutExpired:
        p.kill(); expect(f"{name} : se ferme", False)

log = Path.home() / "Library" / "Logs" / "FileForge.log"
txt = log.read_text(encoding="utf-8", errors="replace") if log.exists() else ""
(OUT / "FileForge-mac.log").write_text(txt, encoding="utf-8")
expect("journal sans erreur Python", "Traceback" not in txt)
print(f"\n{len(fails)} échec(s) : {fails}")
sys.exit(1 if fails else 0)
