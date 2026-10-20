"""Tests propres à Windows, exécutés sur la copie installée de FileForge (voir .github/workflows/windows.yml).

    python windows_ci.py <dossier de résultats>
Conversions Office avec des noms accentués et des espaces, encodages, journal, recherche de LibreOffice.
"""
import os, sys, shutil, traceback, importlib.util as u
from pathlib import Path

D = Path(os.environ["LOCALAPPDATA"]) / "FileForge"
spec = u.spec_from_file_location("ff", D / "fileforge.py")
ff = u.module_from_spec(spec); spec.loader.exec_module(ff)
ff.notify = lambda m: None
ff.choose = lambda items, prompt, default=None, **k: items[0]
T = Path(os.environ.get("TEMP", ".")) / "ff_win_test"
shutil.rmtree(T, ignore_errors=True); T.mkdir()
ok = fail = 0


def check(name, fn):
    global ok, fail
    try:
        r = fn()
        print(f"OK   {name}  {r if r is not None else ''}")
        ok += 1
    except Exception as e:
        print(f"FAIL {name} : {e}")
        traceback.print_exc(limit=3)
        fail += 1


def pages(p):
    import pymupdf
    n = pymupdf.open(p).page_count
    assert n > 0, "PDF vide"
    return f"{Path(p).name} : {n} page(s)"


# fichiers « réels » : noms accentués, espaces, comme celui qui échouait
from pptx import Presentation
from pptx.util import Inches
pr = Presentation()
for i in range(3):
    s = pr.slides.add_slide(pr.slide_layouts[1])
    s.shapes.title.text = f"Séance {i + 1} — été 2026"
    s.placeholders[1].text = "Accents : é è à ç ô, guillemets « »"
PPTX = T / "PPT_PPGS_Séance 1_2026-27.pptx"
pr.save(PPTX)
import openpyxl
wb = openpyxl.Workbook(); ws = wb.active
ws.append(["Poste", "Montant €"]); ws.append(["Café", 3.5]); ws.append(["Loyer", 700])
XLSX = T / "Budget été 2026.xlsx"
wb.save(XLSX)
(T / "notes Windows.txt").write_bytes("Fichier Bloc-notes : élève, café, à bientôt\r\n".encode("cp1252"))
(T / "export excel.csv").write_bytes("nom;ville\r\nZoé;Orléans\r\nJérôme;Besançon\r\n".encode("cp1252"))

so = ff.soffice()
print("LibreOffice :", so)
check("LibreOffice trouvé", lambda: so or (_ for _ in ()).throw(RuntimeError("introuvable")))
check("PowerPoint → PDF (nom accentué)", lambda: pages(ff.doc_to_pdf([str(PPTX)])[0]))
check("Excel → PDF (nom accentué)", lambda: pages(ff.doc_to_pdf([str(XLSX)])[0]))


def docx_roundtrip():
    src = T / "Lettre à Zoé.docx"
    assert ff.lo_convert(T / "notes Windows.txt", "docx", src), "txt → docx"
    return pages(ff.doc_to_pdf([str(src)])[0])


check("Word → PDF (nom accentué)", docx_roundtrip)
check("PowerPoint → ODP", lambda: Path(ff.office_convert([str(PPTX)])[0]).name)
check("Texte Bloc-notes (cp1252) → PDF", lambda: pages(ff.doc_to_pdf([str(T / "notes Windows.txt")])[0]))


def csv_cp1252():
    ff.choose = lambda items, prompt, default=None, **k: next(x for x in items if "xlsx" in x.lower() or "excel" in x.lower())
    out = ff.ACTIONS[[a["fn"].__name__ for a in ff.ACTIONS].index("tab_convert")]["fn"]([str(T / "export excel.csv")])
    wb = openpyxl.load_workbook(out[0])
    vals = [c.value for row in wb.active.iter_rows() for c in row]
    assert "Zoé" in vals and "Besançon" in vals, vals
    return Path(out[0]).name


check("CSV Excel (cp1252) → XLSX sans accents cassés", csv_cp1252)


def log_utf8():
    ff.log("test journal → 📄 é « »")
    assert "test journal → 📄 é « »" in ff.LOG.read_text(encoding="utf-8")


check("Journal en UTF-8", log_utf8)


def pages_retry():
    import pymupdf
    d = pymupdf.open()
    d.new_page(); d.new_page()
    p = T / "deux pages.pdf"; d.save(p)
    answers = iter(["9", "abc", "2"])
    prompts = []
    ff.ask = lambda q, default="", hidden=False: (prompts.append(q), next(answers))[1]
    out = ff.ACTIONS[[a["fn"].__name__ for a in ff.ACTIONS].index("pdf_extract")]["fn"]([str(p)])
    assert pymupdf.open(out[0]).page_count == 1
    assert len(prompts) == 3 and prompts[1].startswith("⚠"), prompts
    return "saisie invalide → redemandée"


check("Pages hors limites : on redemande", pages_retry)

print(f"\n{ok} OK / {fail} échecs")
sys.exit(1 if fail else 0)
